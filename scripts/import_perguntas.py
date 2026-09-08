#!/usr/bin/env python3
"""
Importa as perguntas oficiais do checklist para uma migration SQL.

O script lê a aba "Perguntas" do modelo XLSX ou as tabelas oficiais do DOCX,
valida todo o conteúdo e só então escreve SQL. Ele nunca conecta ao banco.

Uso:
    python scripts/import_perguntas.py --check
    python scripts/import_perguntas.py --fonte perguntas.xlsx --out supabase/migrations/<migration>.sql
    python scripts/import_perguntas.py --fonte perguntas.docx --out supabase/migrations/<migration>.sql
    python scripts/import_perguntas.py --self-check

Qualquer inconsistência faz o processo terminar com código diferente de zero e
nenhum SQL é escrito. As perguntas já existentes são arquivadas pela migration
gerada, preservando respostas de crianças que referenciam seus IDs.
"""

import argparse
import re
import sys
import tempfile
import unicodedata
import zipfile
from collections import Counter
from pathlib import Path

FONTE_PADRAO = Path(
    "AvanceKids-DOCUMENTACAO/LOGICA-ATUALIZADA/"
    "2026.08.18_Logica App para exercícios.docx"
)
ABA = "Perguntas"

COLUNAS = ["Tipo", "Faixa", "Habilidade", "Ordem", "Pergunta"]
KINDS = {"inicial", "triagem"}
FAIXAS = {"F01A", "F02A", "F03A", "F04A", "F05A", "F06A"}

# Rótulo usado no modelo -> skills.key no banco.
HABILIDADE_PARA_KEY = {
    "Comunicação": "comunicacao",
    "Social": "social",
    "Cognitiva": "cognitiva",
    "Coordenação Motora": "motora",
    "Funcional": "funcional",
}

# Código usado nas tabelas do checklist completo -> habilidade do sistema.
SUFIXO_PARA_HABILIDADE = {
    "AC": "Comunicação",
    "AS": "Social",
    "AG": "Cognitiva",
    "AM": "Coordenação Motora",
    "AF": "Funcional",
}

# Os códigos AT são os pré-requisitos. Eles não codificam a habilidade, então
# o vínculo vem da subárea escrita no próprio documento. O skill_id não altera
# a regra de rebaixamento; ele mantém a origem pedagógica da pergunta.
SUBAREA_INICIAL_PARA_HABILIDADE = {
    "contato visual": "Comunicação",
    "comunicacao funcional": "Comunicação",
    "linguagem receptiva": "Comunicação",
    "motora global": "Coordenação Motora",
    "imitacao": "Social",
    "brincar": "Social",
    "brincar com regras": "Social",
    "tolerancia": "Social",
    "funcional": "Funcional",
}

MARCA_EXEMPLO = "EXEMPLO — apague esta linha"
EXEMPLOS_SEM_TIPO = {
    (
        "F01A",
        "Comunicação",
        1,
        "A criança olha ou responde quando é chamada pelo nome?",
    ),
    (
        "F01A",
        "Social",
        2,
        "A criança demonstra interesse por outras pessoas e busca interação?",
    ),
}


class ErroDeValidacao(Exception):
    pass


def normalizar_texto(valor: object) -> str:
    return str(valor or "").strip()


def normalizar_busca(valor: object) -> str:
    """Normaliza caixa e acentos para comparar os rótulos do modelo."""
    texto = unicodedata.normalize("NFKD", normalizar_texto(valor))
    return texto.encode("ascii", "ignore").decode().casefold()


HABILIDADES_NORMALIZADAS = {
    normalizar_busca(rotulo): (rotulo, key)
    for rotulo, key in HABILIDADE_PARA_KEY.items()
}
EXEMPLOS_NORMALIZADOS = {
    (
        faixa.upper(),
        normalizar_busca(habilidade),
        ordem,
        normalizar_busca(pergunta),
    )
    for faixa, habilidade, ordem, pergunta in EXEMPLOS_SEM_TIPO
}


def ler_planilha(caminho: Path) -> list[dict]:
    import openpyxl

    try:
        wb = openpyxl.load_workbook(caminho, read_only=True, data_only=True)
    except (OSError, ValueError, zipfile.BadZipFile) as err:
        raise ErroDeValidacao(f"não foi possível abrir {caminho}: {err}") from err

    try:
        if ABA not in wb.sheetnames:
            raise ErroDeValidacao(
                f"aba '{ABA}' não encontrada em {caminho} (abas: {wb.sheetnames})"
            )

        ws = wb[ABA]
        linhas = ws.iter_rows(values_only=True)
        cabecalho_lido = next(linhas, None)
        if cabecalho_lido is None:
            raise ErroDeValidacao(f"a aba '{ABA}' está vazia")

        cabecalho = [normalizar_texto(celula) for celula in cabecalho_lido]
        if cabecalho != COLUNAS:
            raise ErroDeValidacao(
                "cabeçalho diferente do esperado.\n"
                f"  esperado: {COLUNAS}\n"
                f"  lido:     {cabecalho}"
            )

        registros = []
        for numero_linha, linha in enumerate(linhas, start=2):
            if all(celula is None or normalizar_texto(celula) == "" for celula in linha):
                continue
            registro = dict(zip(COLUNAS, linha))
            registro["_linha_xlsx"] = numero_linha
            registros.append(registro)
        return registros
    finally:
        wb.close()


def ler_docx(caminho: Path) -> list[dict]:
    """Extrai somente tabelas de perguntas; ignora anexos e programas ABA."""
    try:
        from docx import Document
        from docx.opc.exceptions import PackageNotFoundError
    except ImportError as err:
        raise ErroDeValidacao(
            "python-docx não está instalado; instale o pacote para ler a fonte .docx"
        ) from err

    try:
        documento = Document(caminho)
    except (OSError, ValueError, PackageNotFoundError) as err:
        raise ErroDeValidacao(f"não foi possível abrir {caminho}: {err}") from err

    registros: list[dict] = []
    linha_logica = 1
    cabecalho_perguntas = [
        "codigo",
        "subarea",
        "habilidade",
        "exemplo para o pai",
    ]

    for numero_tabela, tabela in enumerate(documento.tables, start=1):
        if not tabela.rows:
            continue

        cabecalho = [normalizar_busca(c.text) for c in tabela.rows[0].cells]
        if cabecalho[:4] != cabecalho_perguntas:
            continue

        for numero_linha, linha in enumerate(tabela.rows[1:], start=2):
            valores = [normalizar_texto(c.text) for c in linha.cells]
            if not any(valores):
                continue

            codigo = valores[0].upper()
            match = re.fullmatch(r"(F\d{2})(AT|AC|AS|AG|AM|AF)(\d{3})", codigo)
            if match is None:
                raise ErroDeValidacao(
                    f"tabela {numero_tabela}, linha {numero_linha}: "
                    f"código de pergunta inválido {codigo!r}"
                )

            prefixo_faixa, familia, sufixo_ordem = match.groups()
            faixa = f"{prefixo_faixa}A"
            if familia == "AT":
                kind = "Inicial"
                subarea = normalizar_busca(valores[1])
                habilidade = SUBAREA_INICIAL_PARA_HABILIDADE.get(subarea)
                if habilidade is None:
                    raise ErroDeValidacao(
                        f"tabela {numero_tabela}, linha {numero_linha}: "
                        f"subárea inicial desconhecida {valores[1]!r}"
                    )
            else:
                kind = "Triagem"
                habilidade = SUFIXO_PARA_HABILIDADE.get(familia)
                if habilidade is None:
                    raise ErroDeValidacao(
                        f"tabela {numero_tabela}, linha {numero_linha}: "
                        f"família de código desconhecida {familia!r}"
                    )

            linha_logica += 1
            registros.append(
                {
                    "Tipo": kind,
                    "Faixa": faixa,
                    "Habilidade": habilidade,
                    "Ordem": int(sufixo_ordem),
                    "Pergunta": valores[3],
                    "_linha_xlsx": linha_logica,
                }
            )

    if not registros:
        raise ErroDeValidacao(
            "nenhuma tabela com cabeçalho Código/Subárea/Habilidade/Exemplo para o Pai foi encontrada"
        )
    return registros


def ler_fonte(caminho: Path) -> list[dict]:
    sufixo = caminho.suffix.casefold()
    if sufixo == ".xlsx":
        return ler_planilha(caminho)
    if sufixo == ".docx":
        return ler_docx(caminho)
    raise ErroDeValidacao(
        f"formato {caminho.suffix or '(sem extensão)'} não suportado; use .xlsx ou .docx"
    )


def converter_ordem(valor: object) -> int | None:
    if isinstance(valor, bool):
        return None
    if isinstance(valor, int):
        return valor
    if isinstance(valor, float):
        return int(valor) if valor.is_integer() else None

    texto = normalizar_texto(valor)
    if re.fullmatch(r"[+-]?\d+", texto):
        return int(texto)
    return None


def eh_linha_de_exemplo(registro: dict) -> bool:
    if normalizar_busca(registro.get("Tipo")) == normalizar_busca(MARCA_EXEMPLO):
        return True

    ordem = converter_ordem(registro.get("Ordem"))
    assinatura = (
        normalizar_texto(registro.get("Faixa")).upper(),
        normalizar_busca(registro.get("Habilidade")),
        ordem,
        normalizar_busca(registro.get("Pergunta")),
    )
    return assinatura in EXEMPLOS_NORMALIZADOS


def validar(registros: list[dict]) -> tuple[list[dict], list[str]]:
    """Devolve (perguntas válidas, problemas), sem interromper no primeiro erro."""
    perguntas: list[dict] = []
    problemas: list[str] = []
    vistos: dict[tuple[str, str, str, int], int] = {}

    if not registros:
        problemas.append("nenhuma pergunta encontrada na aba 'Perguntas'")
        return perguntas, problemas

    habilidades_aceitas = ", ".join(HABILIDADE_PARA_KEY)

    for indice, registro in enumerate(registros, start=2):
        linha = int(registro.get("_linha_xlsx", indice))

        if eh_linha_de_exemplo(registro):
            problemas.append(
                f"linha {linha}: linha de exemplo do modelo; apague-a antes de importar"
            )
            continue

        erros_da_linha: list[str] = []

        kind = normalizar_texto(registro.get("Tipo")).casefold()
        if kind not in KINDS:
            erros_da_linha.append(
                f"tipo desconhecido {registro.get('Tipo')!r}; use Inicial ou Triagem"
            )

        faixa = normalizar_texto(registro.get("Faixa"))
        if faixa not in FAIXAS:
            erros_da_linha.append(
                f"faixa {registro.get('Faixa')!r} não existe em age_brackets; "
                f"use um de: {', '.join(sorted(FAIXAS))}"
            )

        habilidade_lida = normalizar_texto(registro.get("Habilidade"))
        habilidade = HABILIDADES_NORMALIZADAS.get(normalizar_busca(habilidade_lida))
        if habilidade is None:
            erros_da_linha.append(
                f"habilidade {registro.get('Habilidade')!r} não corresponde a skills.key; "
                f"use uma de: {habilidades_aceitas}"
            )

        ordem = converter_ordem(registro.get("Ordem"))
        if ordem is None:
            erros_da_linha.append(
                f"ordem {registro.get('Ordem')!r} não é um número inteiro"
            )

        texto = normalizar_texto(registro.get("Pergunta"))
        if not texto:
            erros_da_linha.append("pergunta vazia")

        if erros_da_linha:
            problemas.extend(f"linha {linha}: {erro}" for erro in erros_da_linha)
            continue

        assert habilidade is not None and ordem is not None
        rotulo_habilidade, skill_key = habilidade
        chave = (kind, skill_key, faixa, ordem)
        linha_anterior = vistos.get(chave)
        if linha_anterior is not None:
            problemas.append(
                f"linha {linha}: ordem {ordem} duplicada em "
                f"({kind}, {rotulo_habilidade}, {faixa}); primeira ocorrência na "
                f"linha {linha_anterior}"
            )
            continue
        vistos[chave] = linha

        perguntas.append(
            {
                "kind": kind,
                "faixa": faixa,
                "skill_key": skill_key,
                "habilidade": rotulo_habilidade,
                "ordem": ordem,
                "texto": texto,
            }
        )

    return perguntas, problemas


def resumo(perguntas: list[dict]) -> str:
    por_kind = Counter(item["kind"] for item in perguntas)
    por_faixa = Counter(item["faixa"] for item in perguntas)
    por_skill = Counter(item["skill_key"] for item in perguntas)
    return "\n".join(
        [
            f"perguntas válidas: {len(perguntas)}",
            f"  por tipo:       {dict(sorted(por_kind.items()))}",
            f"  por faixa:      {dict(sorted(por_faixa.items()))}",
            f"  por habilidade: {dict(sorted(por_skill.items()))}",
        ]
    )


def sql_literal(valor: object) -> str:
    return "'" + str(valor).replace("'", "''") + "'"


def sql_array(valores: list[str]) -> str:
    return "ARRAY[" + ", ".join(sql_literal(valor) for valor in valores) + "]::text[]"


def gerar_sql(perguntas: list[dict], fonte: Path) -> str:
    """Gera SQL determinístico; só deve ser chamado depois de validar sem erros."""
    linhas = sorted(
        perguntas,
        key=lambda item: (
            item["kind"],
            item["faixa"],
            item["skill_key"],
            item["ordem"],
            item["texto"],
        ),
    )
    if not linhas:
        raise ErroDeValidacao("não é permitido gerar uma migration sem perguntas")

    valores = [
        "    ("
        + ", ".join(
            [
                sql_literal(item["kind"]),
                sql_literal(item["skill_key"]),
                sql_literal(item["faixa"]),
                str(item["ordem"]),
                sql_literal(item["texto"]),
            ]
        )
        + ")"
        for item in linhas
    ]

    faixas = sorted({item["faixa"] for item in linhas})
    skills = sorted({item["skill_key"] for item in linhas})
    total = len(linhas)

    return f"""-- ============================================================
-- Perguntas oficiais do checklist
--
-- GERADO POR scripts/import_perguntas.py — NÃO EDITAR À MÃO.
-- Fonte: {fonte.as_posix()}
--        ({total} perguntas oficiais validadas)
--
-- As perguntas anteriores são arquivadas, não apagadas: respostas já
-- registradas continuam apontando para os mesmos IDs.
-- ============================================================

-- Confere as referências antes de alterar qualquer pergunta. Uma diferença
-- entre a planilha e o banco aborta a migration inteira.
DO $$
DECLARE
  v_faixas INTEGER;
  v_skills INTEGER;
BEGIN
  SELECT count(*) INTO v_faixas
  FROM age_brackets
  WHERE codigo = ANY({sql_array(faixas)});

  IF v_faixas <> {len(faixas)} THEN
    RAISE EXCEPTION 'esperava {len(faixas)} faixas da planilha em age_brackets, encontrei %', v_faixas;
  END IF;

  SELECT count(*) INTO v_skills
  FROM skills
  WHERE key = ANY({sql_array(skills)});

  IF v_skills <> {len(skills)} THEN
    RAISE EXCEPTION 'esperava {len(skills)} habilidades da planilha em skills, encontrei %', v_skills;
  END IF;
END $$;

-- Preserva as perguntas e as respostas antigas, mas deixa somente o conteúdo
-- oficial novo visível para próximas avaliações.
UPDATE questions
SET status = 'arquivado'
WHERE status = 'ativo';

WITH oficial(kind, skill_key, faixa_codigo, ordem, texto) AS (
  VALUES
{',\n'.join(valores)}
)
INSERT INTO questions (kind, skill_id, age_bracket_id, texto, ordem, status)
SELECT
  o.kind::question_kind,
  s.id,
  b.id,
  o.texto,
  o.ordem,
  'ativo'::record_status
FROM oficial o
JOIN skills s ON s.key = o.skill_key
JOIN age_brackets b ON b.codigo = o.faixa_codigo;

-- Se algum JOIN perder uma linha, a exceção faz a migration falhar em vez de
-- deixar um checklist incompleto.
DO $$
DECLARE
  v_total INTEGER;
BEGIN
  SELECT count(*) INTO v_total
  FROM questions
  WHERE status = 'ativo';

  IF v_total <> {total} THEN
    RAISE EXCEPTION 'esperava {total} perguntas oficiais ativas, encontrei %', v_total;
  END IF;
END $$;
"""


def criar_xlsx_sintetico(caminho: Path, linhas: list[list[object]]) -> None:
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.title = ABA
    ws.append(COLUNAS)
    for linha in linhas:
        ws.append(linha)
    wb.save(caminho)
    wb.close()


def self_check() -> None:
    """Exercita leitura e validações usando apenas planilhas sintéticas."""
    linhas_validas = [
        ["Inicial", "F01A", "Comunicação", 1, "Pergunta inicial válida?"],
        ["Triagem", "F02A", "Social", 1, "Pergunta d'água válida?"],
    ]

    def conferir_problema(
        pasta: Path, nome: str, linhas: list[list[object]], trecho: str
    ) -> None:
        caminho = pasta / f"{nome}.xlsx"
        criar_xlsx_sintetico(caminho, linhas)
        _, problemas = validar(ler_planilha(caminho))
        assert any(trecho in problema for problema in problemas), problemas

    with tempfile.TemporaryDirectory(prefix="import_perguntas_") as tmp:
        pasta = Path(tmp)

        caminho_valido = pasta / "valido.xlsx"
        criar_xlsx_sintetico(caminho_valido, linhas_validas)
        perguntas, problemas = validar(ler_planilha(caminho_valido))
        assert problemas == [], problemas
        assert len(perguntas) == 2
        assert perguntas[0]["skill_key"] == "comunicacao"
        assert perguntas[1]["kind"] == "triagem"

        conferir_problema(
            pasta,
            "faixa_inexistente",
            [["Inicial", "F07A", "Comunicação", 1, "Pergunta?"]],
            "não existe em age_brackets",
        )
        conferir_problema(
            pasta,
            "habilidade_inexistente",
            [["Inicial", "F01A", "Linguagem", 1, "Pergunta?"]],
            "não corresponde a skills.key",
        )
        conferir_problema(
            pasta,
            "kind_invalido",
            [["Diagnóstico", "F01A", "Comunicação", 1, "Pergunta?"]],
            "tipo desconhecido",
        )
        conferir_problema(
            pasta,
            "ordem_duplicada",
            [linhas_validas[0], ["Inicial", "F01A", "Comunicação", 1, "Outra?"]],
            "ordem 1 duplicada",
        )
        conferir_problema(
            pasta,
            "linha_exemplo",
            [[MARCA_EXEMPLO, "F01A", "Comunicação", 1, "Pergunta?"]],
            "linha de exemplo do modelo",
        )
        conferir_problema(
            pasta,
            "conteudo_exemplo",
            [["Inicial", *next(iter(EXEMPLOS_SEM_TIPO))]],
            "linha de exemplo do modelo",
        )

        sql = gerar_sql(perguntas, Path("perguntas-sinteticas.xlsx"))
        assert "INSERT INTO questions" in sql
        assert "UPDATE questions" in sql and "status = 'arquivado'" in sql
        assert "Pergunta d''água válida?" in sql
        assert "esperava 2 perguntas oficiais ativas" in sql
        assert "('inicial', 'comunicacao', 'F01A', 1, 'Pergunta inicial válida?')" in sql

    print("self-check: OK")


def main() -> int:
    parser = argparse.ArgumentParser(
        description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter
    )
    parser.add_argument(
        "--fonte", type=Path, default=FONTE_PADRAO, help="caminho do .xlsx ou .docx"
    )
    parser.add_argument("--out", help="arquivo .sql a gerar")
    parser.add_argument("--check", action="store_true", help="apenas valida e resume")
    parser.add_argument("--self-check", action="store_true", help="roda os testes do script")
    args = parser.parse_args()

    if args.self_check:
        self_check()
        return 0

    fonte = args.fonte
    if not fonte.exists():
        print(f"ERRO: arquivo fonte não encontrado: {fonte}", file=sys.stderr)
        return 2

    try:
        registros = ler_fonte(fonte)
    except ErroDeValidacao as err:
        print(f"ERRO: {err}", file=sys.stderr)
        return 2

    perguntas, problemas = validar(registros)
    print(resumo(perguntas))

    if problemas:
        print(f"\n{len(problemas)} PROBLEMA(S) — nada foi gerado:", file=sys.stderr)
        for problema in problemas:
            print(f"  - {problema}", file=sys.stderr)
        return 1

    print("\nvalidação: OK")

    if args.check or not args.out:
        return 0

    try:
        conteudo = gerar_sql(perguntas, fonte)
    except ErroDeValidacao as err:
        print(f"ERRO: {err}", file=sys.stderr)
        return 2

    destino = Path(args.out)
    destino.parent.mkdir(parents=True, exist_ok=True)
    destino.write_text(conteudo, encoding="utf-8", newline="\n")
    print(f"gerado: {destino} ({destino.stat().st_size} bytes)")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
