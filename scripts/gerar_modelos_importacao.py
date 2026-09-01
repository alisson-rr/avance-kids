#!/usr/bin/env python3
"""
Gera a planilha-modelo que a cliente preenche para importar conteúdo.

    python scripts/gerar_modelos_importacao.py              # escreve o .xlsx
    python scripts/gerar_modelos_importacao.py --self-check # só valida, não escreve

Saída: docs/modelos-importacao.xlsx, com três abas —
  * Instruções: como preencher, em linguagem de quem vai preencher.
  * Perguntas:  as perguntas do checklist (tabela `questions`).
  * Atividades: os programas ABA (tabela `exercises`), no MESMO formato de 17
                colunas que já importou em migration-08. O cabeçalho é lido de
                `import_programas.py` para não existirem duas listas.

O que NÃO entra na planilha, de propósito: imagem/vídeo, marcação de
Gratuito/Premium e arquivar/reativar. Os três já são campos do cadastro de
atividade no backoffice, e mídia envolve upload de arquivo — coisa que planilha
não carrega. Duplicar isso aqui criaria duas fontes de verdade para o mesmo
dado.

O --self-check compara as faixas etárias com a asserção de
scripts/validate_migrations.sh. É o que impede o modelo de envelhecer calado
quando os limites mudarem de novo.
"""

import argparse
import re
import sys
from pathlib import Path

RAIZ = Path(__file__).resolve().parent.parent
SAIDA_PADRAO = RAIZ / "docs" / "modelos-importacao.xlsx"
VALIDATE_SH = RAIZ / "scripts" / "validate_migrations.sh"

sys.path.insert(0, str(Path(__file__).resolve().parent))
from import_programas import COLUNAS as COLUNAS_ATIVIDADES_FONTE  # noqa: E402

# Faixas conforme migration-11. (codigo, rótulo, meses_min, meses_max)
FAIXAS = [
    ("F01A", "12 a 24 meses", 12, 24),
    ("F02A", "25 a 36 meses", 25, 36),
    ("F03A", "37 a 48 meses", 37, 48),
    ("F04A", "49 a 60 meses", 49, 60),
    ("F05A", "5 a 7 anos", 61, 95),
    ("F06A", "8 a 11 anos", 96, 143),
]

# Espelha `skills` no banco e HabilidadeKey no app. A letra é a que aparece no
# código da atividade (F01A**C**001 = Comunicação).
HABILIDADES = [
    ("Comunicação", "C"),
    ("Social", "S"),
    ("Cognitiva", "G"),
    ("Coordenação Motora", "M"),
    ("Funcional", "F"),
]

NIVEIS = ["Aquisição", "Generalização", "Manutenção"]
TIPOS_PERGUNTA = ["Inicial", "Triagem"]

COLUNAS_PERGUNTAS = [
    "Tipo",
    "Faixa",
    "Habilidade",
    "Ordem",
    "Pergunta",
]

COLUNAS_ATIVIDADES = [nome_xlsx for nome_xlsx, _ in COLUNAS_ATIVIDADES_FONTE]

MARCA_EXEMPLO = "EXEMPLO — apague esta linha"

EXEMPLOS_PERGUNTAS = [
    [MARCA_EXEMPLO, "F01A", "Comunicação", 1,
     "A criança olha ou responde quando é chamada pelo nome?"],
    [MARCA_EXEMPLO, "F01A", "Social", 2,
     "A criança demonstra interesse por outras pessoas e busca interação?"],
]

# Uma linha por nível do mesmo código, que é como a planilha oficial já vem.
EXEMPLOS_ATIVIDADES = [
    [MARCA_EXEMPLO, "Imitar sons e palavras simples", "Imitação vocal",
     "Atenção conjunta", nivel,
     "Objetivo do programa.", "Materiais necessários.", "Recursos extras.",
     "Exemplos de brincadeiras.", "Frequência sugerida.",
     "Hierarquia de dicas.", "Passo a passo do procedimento.",
     "Resposta esperada da criança.", "O que fazer quando erra.",
     "Critério para avançar.", "Como registrar os dados.",
     "Exemplos de reforços."]
    for nivel in NIVEIS
]

INSTRUCOES = [
    ("Como preencher esta planilha", None),
    ("", None),
    ("Há duas abas para preencher: Perguntas e Atividades. Preencha só o que for enviar —"
     " uma aba vazia é ignorada.", None),
    ("Apague as linhas de exemplo antes de enviar. Elas estão marcadas na primeira coluna.", None),
    ("Não mude, não traduza e não reordene os títulos das colunas: a importação"
     " confere um por um e para se algum estiver diferente.", None),
    ("", None),
    ("Aba PERGUNTAS", "titulo"),
    ("Tipo: 'Inicial' para as perguntas de pré-requisito da faixa; 'Triagem' para as"
     " perguntas de cada habilidade.", None),
    ("Faixa: o código da faixa etária (ver tabela abaixo).", None),
    ("Habilidade: uma das cinco. Toda pergunta precisa de uma, inclusive as do tipo Inicial.", None),
    ("Ordem: a sequência em que a pergunta aparece. Números inteiros, sem repetir dentro"
     " da mesma combinação de Tipo + Faixa + Habilidade.", None),
    ("Pergunta: o texto exato que o responsável vai ler.", None),
    ("", None),
    ("Importante: as perguntas do tipo Inicial são as que decidem se a criança começa numa"
     " faixa mais baixa — duas respostas 'quase nunca' rebaixam uma faixa. Por isso vale"
     " manter a mesma quantidade de perguntas Inicial em todas as faixas.", None),
    ("", None),
    ("Aba ATIVIDADES", "titulo"),
    ("É o mesmo formato de 17 colunas da planilha que você já enviou. Se for atualizar"
     " aquele arquivo, pode enviá-lo direto: não precisa recopiar para cá.", None),
    ("Cada código precisa de exatamente 3 linhas, uma para cada Nível"
     " (Aquisição, Generalização, Manutenção).", None),
    ("Código: a faixa, a letra da habilidade e o número sequencial, sem espaços —"
     " por exemplo F01AC001. A letra de cada habilidade está na tabela abaixo.", None),
    ("Nenhuma célula pode ficar vazia nas 17 colunas.", None),
    ("", None),
    ("O que NÃO vai nesta planilha", "titulo"),
    ("Imagem ou vídeo da atividade, marcação de Gratuito/Premium e arquivar/reativar são"
     " definidos no painel administrativo, atividade por atividade. Ficam de fora daqui"
     " para não haver dois lugares dizendo coisas diferentes sobre o mesmo item.", None),
]


def carregar_faixas_do_validador() -> list[tuple[str, int, int]]:
    """Lê os limites que scripts/validate_migrations.sh trava, para comparar."""
    texto = VALIDATE_SH.read_text(encoding="utf-8")
    linha = next(
        (l for l in texto.splitlines() if "('F01A'," in l and "('F06A'," in l),
        None,
    )
    if linha is None:
        raise SystemExit(
            "não achei a asserção das 6 faixas em scripts/validate_migrations.sh — "
            "o modelo não tem como se conferir sozinho"
        )
    achados = re.findall(r"\('(F0\dA)',(\d+),(\d+)\)", linha)
    return [(cod, int(mn), int(mx)) for cod, mn, mx in achados]


def self_check() -> None:
    esperado = [(cod, mn, mx) for cod, _, mn, mx in FAIXAS]
    lido = carregar_faixas_do_validador()
    if lido != esperado:
        raise SystemExit(
            "faixas do modelo divergem de scripts/validate_migrations.sh:\n"
            f"  modelo:    {esperado}\n"
            f"  validador: {lido}"
        )

    # Contíguas, como a migration-11 garante no banco.
    for (_, _, _, max_anterior), (cod, _, min_atual, _) in zip(FAIXAS, FAIXAS[1:]):
        if min_atual != max_anterior + 1:
            raise SystemExit(f"lacuna antes de {cod}: {max_anterior} -> {min_atual}")

    if len(COLUNAS_ATIVIDADES) != 17:
        raise SystemExit(f"esperava 17 colunas de atividade, import_programas tem {len(COLUNAS_ATIVIDADES)}")

    for linha in EXEMPLOS_ATIVIDADES:
        if len(linha) != len(COLUNAS_ATIVIDADES):
            raise SystemExit("linha de exemplo com número de colunas diferente do cabeçalho")
    for linha in EXEMPLOS_PERGUNTAS:
        if len(linha) != len(COLUNAS_PERGUNTAS):
            raise SystemExit("exemplo de pergunta com número de colunas diferente do cabeçalho")

    # A marca de exemplo tem de ser recusada pelo importador (o código não casa
    # com o padrão), senão uma linha esquecida entraria como conteúdo real.
    from import_programas import CODIGO_RE

    if CODIGO_RE.match(MARCA_EXEMPLO):
        raise SystemExit("a marca de exemplo passaria pela validação de código")

    print("self-check ok: faixas, colunas e linhas de exemplo conferem")


def gerar(saida: Path) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.datavalidation import DataValidation

    negrito = Font(bold=True)
    fundo = PatternFill("solid", fgColor="DDE8FF")
    quebra = Alignment(wrap_text=True, vertical="top")

    wb = Workbook()

    # ── Instruções ──
    ws = wb.active
    ws.title = "Instruções"
    ws.column_dimensions["A"].width = 110
    for i, (texto, estilo) in enumerate(INSTRUCOES, start=1):
        celula = ws.cell(row=i, column=1, value=texto)
        celula.alignment = quebra
        if i == 1 or estilo == "titulo":
            celula.font = negrito

    linha = len(INSTRUCOES) + 2
    ws.cell(row=linha, column=1, value="Faixas etárias").font = negrito
    linha += 1
    for cod, rotulo, mn, mx in FAIXAS:
        # Os rótulos das quatro primeiras já são em meses; repetir ficaria bobo.
        faixa_em_meses = f" ({mn} a {mx} meses)" if "meses" not in rotulo else ""
        ws.cell(row=linha, column=1, value=f"{cod} — {rotulo}{faixa_em_meses}")
        linha += 1

    linha += 1
    ws.cell(row=linha, column=1, value="Habilidades e a letra usada no código").font = negrito
    linha += 1
    for rotulo, letra in HABILIDADES:
        ws.cell(row=linha, column=1, value=f"{letra} — {rotulo}")
        linha += 1

    def montar_aba(nome: str, colunas: list[str], exemplos: list[list], larguras: dict[int, int]):
        aba = wb.create_sheet(nome)
        for c, titulo in enumerate(colunas, start=1):
            celula = aba.cell(row=1, column=c, value=titulo)
            celula.font = negrito
            celula.fill = fundo
            celula.alignment = quebra
        for r, linha_exemplo in enumerate(exemplos, start=2):
            for c, valor in enumerate(linha_exemplo, start=1):
                aba.cell(row=r, column=c, value=valor).alignment = quebra
        for c in range(1, len(colunas) + 1):
            aba.column_dimensions[get_column_letter(c)].width = larguras.get(c, 34)
        aba.freeze_panes = "A2"
        return aba

    # ── Perguntas ──
    perguntas = montar_aba(
        "Perguntas", COLUNAS_PERGUNTAS, EXEMPLOS_PERGUNTAS,
        {1: 26, 2: 10, 3: 22, 4: 8, 5: 80},
    )
    listas = [
        (TIPOS_PERGUNTA, "A"),
        ([c for c, _, _, _ in FAIXAS], "B"),
        ([r for r, _ in HABILIDADES], "C"),
    ]
    for valores, coluna in listas:
        dv = DataValidation(type="list", formula1='"' + ",".join(valores) + '"', allow_blank=True)
        perguntas.add_data_validation(dv)
        dv.add(f"{coluna}2:{coluna}1000")

    # ── Atividades ──
    atividades = montar_aba(
        "Atividades", COLUNAS_ATIVIDADES, EXEMPLOS_ATIVIDADES,
        {1: 26, 2: 34, 3: 26, 4: 22, 5: 16},
    )
    dv_nivel = DataValidation(type="list", formula1='"' + ",".join(NIVEIS) + '"', allow_blank=True)
    atividades.add_data_validation(dv_nivel)
    dv_nivel.add("E2:E2000")

    saida.parent.mkdir(parents=True, exist_ok=True)
    wb.save(saida)
    print(f"modelo escrito em {saida.relative_to(RAIZ)}")
    print(f"  Perguntas:  {len(COLUNAS_PERGUNTAS)} colunas")
    print(f"  Atividades: {len(COLUNAS_ATIVIDADES)} colunas (mesmo formato de migration-08)")


def main() -> int:
    ap = argparse.ArgumentParser(description=__doc__)
    ap.add_argument("--out", type=Path, default=SAIDA_PADRAO)
    ap.add_argument("--self-check", action="store_true", help="valida e sai, sem escrever")
    args = ap.parse_args()

    self_check()
    if args.self_check:
        return 0
    gerar(args.out)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
